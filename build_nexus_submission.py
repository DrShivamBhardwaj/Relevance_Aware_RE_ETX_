from pathlib import Path
import re, shutil, zipfile, os, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "submission" / "nexus"
LATEX = OUT / "latex"
OUT.mkdir(parents=True, exist_ok=True)
LATEX.mkdir(parents=True, exist_ok=True)
(LATEX / "figures").mkdir(parents=True, exist_ok=True)

TITLE = "Energy-Information Co-Design for Resource-Bounded Edge Intelligence: A Systems-Learning Federated Framework"
HIGHLIGHTS = [
    "Network state is exposed directly to federated learning orchestration",
    "Control-plane signaling is measured rather than assumed negligible",
    "Matched controls reveal Pareto trade-offs between accuracy and relay cost",
    "ns-3 replay links offered-load reduction to packet-network behavior",
]
IN_BRIEF = ("Tripathi et al. frame federated edge intelligence as an energy-information co-design problem. "
            "Their controller coordinates utility-target participation, update fidelity, relay pressure, and "
            "staleness while explicitly pricing pre-selection signaling. Held-out real-data experiments and "
            "ns-3 low-power wireless replay expose workload-dependent Pareto trade-offs rather than universal dominance.")
BROADER = (
    "Resource-bounded edge intelligence increasingly supports environmental sensing, infrastructure observation, "
    "wearable analytics, and industrial monitoring. In these systems, selecting a learning client is also a physical "
    "resource-allocation decision because the selected update determines network load, relay burden, and delay. "
    "This study develops a systems-learning co-design framework that exposes multi-hop network state directly to "
    "federated orchestration. The evaluation deliberately separates learning gains from network costs, prices control "
    "metadata instead of treating scheduler information as free, and distinguishes open-loop packet replay from "
    "closed-loop co-simulation. The resulting evidence is best interpreted as a configurable Pareto mechanism for "
    "resource-bounded distributed AI rather than a universal efficiency or sustainability claim."
)


FIGURES = [
    ("figures/final/02_system_architecture.png", "Figure 1. Conceptual cross-layer HFL architecture. Learning-capable clients exchange models through multi-hop relays and edge gateways. Downlink dissemination is shown for system completeness; reported communication totals in this study refer to uplink model-update traffic unless stated otherwise."),
    ("figures/final/01_network_topology.png", "Figure 2. Intel Berkeley Lab WSN measured topology abstraction. Gateway motes are the configured edge gateways. Dashed links show a visually filtered subset of measured bidirectional-connectivity support; routing uses the full measured connectivity matrix and ETX values."),
    ("figures/final/07_intel_berkeley_results.png", "Figure 3. Intel held-out evaluation at c = 0.9. Raw means are annotated and tabulated. The line plot normalizes each metric to its own within-metric maximum only for visual comparison; it is not a composite performance score. Confidence intervals are reported in Table 2."),
    ("figures/final/06_uci_har_results.png", "Figure 4. UCI HAR held-out evaluation at c = 0.9 using the blocked overlap-safe within-client split. Raw means are annotated and tabulated. The line plot uses within-metric normalization only for visualization; it is not a composite score. Confidence intervals are reported in Table 3."),
    ("figures/final/04_ablation.png", "Figure 5. Held-out component ablations on Intel and UCI HAR at c = 0.9. 'No deficit' removes the utility-target participation-deficit term. Adaptive fidelity is the principal source of traffic reduction, while relay pressure and participation deficit affect different trade-off dimensions."),
    ("figures/final/03_ns3_validation.png", "Figure 6. ns-3.47 IEEE 802.15.4/LR-WPAN hop-equivalent replay of held-out synthetic traffic. Report delivery ratio, delivered-report delay, and channel-access failures are shown for dense-nominal and 24-sensor-nominal conditions. The HFL controller itself is not executed inside ns-3."),
    ("figures/final/05_learning_communication_tradeoff.png", "Figure 7. Held-out learning-versus-uplink-model-update-traffic trade-offs at c = 0.9. Marker area scales with maximum relay energy. The figure includes compression-matched controls, the FedCG-adapted comparator, and the proposed controller."),
]

ALGORITHM1_ROWS = [
    ("1", "Update current routes, ETX burdens, residual-energy state, and candidate availability."),
    ("2", "Collect one 21-byte control packet per candidate: header, local loss, residual-energy estimate, and availability."),
    ("3", "Compute pre-selection utility U_i(t) and utility-target share pi_i(t)."),
    ("4", "Construct the eligible set using availability and the residual-energy threshold."),
    ("5", "For each eligible client, evaluate every retained Top-k fraction rho in {0.15, 0.30, 0.50, 0.75, 1.00}."),
    ("6", "Choose rho_i(t) that minimizes the implemented route-scarcity-relay-pressure plus distortion surrogate."),
    ("7", "Compute scheduling score Gamma_i(t) and select the K_t highest-scoring eligible clients."),
    ("8", "Update the client participation-deficit queues Q_i(t) from target and realized participation."),
    ("9", "Selected clients perform local optimization and update their utility history from novelty and non-negative local improvement."),
    ("10", "Apply error-feedback Top-k compression and transmit the sparse update over the selected multi-hop route."),
    ("11", "Accumulate source/relay energy and update relay-pressure queues Z_r(t)."),
    ("12", "Admit updates whose modeled arrival time falls within the staleness limit; discard updates that exceed it."),
    ("13", "Aggregate arrived client updates at each edge using sample count, staleness decay, and update utility."),
    ("14", "Normalize edge contributions at the cloud, update the global model, and record true two-level cloud influence."),
]

def transform_source():
    src = (ROOT / "MANUSCRIPT_RECONSTRUCTION.md").read_text().strip()
    front = "# " + TITLE + "\n\n"
    front += "**Abhinandan Tripathi¹, Vijay Kumar Tiwari², Mohd. Arif³, Abhishek Kumar Pandey⁴*, Shivam Bhardwaj⁵**\n\n"
    front += "¹Department of Computer Science and Engineering, Buddha Institute of Technology, Gorakhpur, India\n\n"
    front += "²Department of Information Technology, Madan Mohan Malaviya University of Technology (MMMUT), Gorakhpur, India\n\n"
    front += "³Department of Computer Science and Engineering, Galgotias University, Greater Noida, India\n\n"
    front += "⁴Department of Computer Science and Engineering, Bennett University, Greater Noida, India\n\n"
    front += "⁵United Institute of Management, Prayagraj, India\n\n"
    front += "Emails: ¹abhinandan282@bit.ac.in; ²vktitca@mmmut.ac.in; ³md.arif@galgotiasuniversity.edu.in; ⁴abhishek.pandey2@bennett.edu.in; ⁵shibambhardwaj@gmail.com\n\n"
    front += "ORCID: Abhishek Kumar Pandey — 0000-0003-3799-9754; Shivam Bhardwaj — 0009-0005-4554-7397\n\n"
    front += "*Corresponding author: Abhishek Kumar Pandey (abhishek.pandey2@bennett.edu.in)\n\n"
    front += "All authors contributed equally to this work.\n\n"
    front += "## Highlights\n\n" + "\n".join("- " + h for h in HIGHLIGHTS) + "\n\n"
    front += "## In brief\n\n" + IN_BRIEF + "\n\n## Broader context\n\n" + BROADER + "\n\n"
    return front + src + "\n"

SOURCE = transform_source()
(OUT / "NEXUS_MANUSCRIPT_SOURCE.md").write_text(SOURCE)
OMML_MAP = json.loads((OUT / "omml_math_map.json").read_text())
OMML_INLINE = OMML_MAP["inline"]
OMML_BLOCK = OMML_MAP["block"]

def _append_omml(parent, xml_text):
    parent.append(parse_xml(xml_text))

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def set_cell_text(cell, text, bold=False, color=None, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text); r.bold = bold
    r.font.name = "Times New Roman"; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
    b = OxmlElement("w:instrText"); b.set(qn("xml:space"), "preserve"); b.text = "PAGE"
    c = OxmlElement("w:fldChar"); c.set(qn("w:fldCharType"), "end")
    run._r.append(a); run._r.append(b); run._r.append(c)

def configure_doc(doc):
    sec = doc.sections[0]
    sec.top_margin=Inches(.75); sec.bottom_margin=Inches(.75); sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
    n=doc.styles["Normal"]; n.font.name="Times New Roman"; n.font.size=Pt(11); n.paragraph_format.space_after=Pt(5); n.paragraph_format.line_spacing=1.08; n.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for sname,size in [("Title",16),("Heading 1",13),("Heading 2",11.5),("Heading 3",11)]:
        st=doc.styles[sname]; st.font.name="Times New Roman"; st.font.size=Pt(size); st.font.bold=True
    add_page_number(sec.footer.paragraphs[0])

def parse_inline_runs(p, text):
    pos=0
    pat=re.compile(r"(\*\*.*?\*\*|\*[^*]+\*|\\\(.*?\\\))")
    for m in pat.finditer(text):
        if m.start()>pos:
            rr=p.add_run(text[pos:m.start()]); rr.font.name="Times New Roman"
        token=m.group(0)
        if token.startswith("\\("):
            key=token[2:-2].strip()
            if key not in OMML_INLINE:
                raise KeyError(f"Missing inline OMML mapping: {key}")
            _append_omml(p._p, OMML_INLINE[key])
        else:
            rr=p.add_run(token.strip("*")); rr.font.name="Times New Roman"
            if token.startswith("**"): rr.bold=True
            else: rr.italic=True
        pos=m.end()
    if pos<len(text):
        rr=p.add_run(text[pos:]); rr.font.name="Times New Roman"

def add_math_para(doc, text):
    key=" ".join(text.split())
    if key not in OMML_BLOCK:
        raise KeyError(f"Missing block OMML mapping: {key}")
    p=doc.add_paragraph()
    _append_omml(p._p, OMML_BLOCK[key])
    return p

def add_table_caption(doc, text):
    clean=re.sub(r"^\*\*|\*\*$", "", text).strip()
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(6)
    p.paragraph_format.space_after=Pt(3)
    r=p.add_run(clean)
    r.bold=True
    r.font.name="Times New Roman"
    r.font.size=Pt(9)
    return p

def add_algorithm_box(doc):
    cap=doc.add_paragraph()
    cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before=Pt(6)
    cap.paragraph_format.space_after=Pt(3)
    rr=cap.add_run("Algorithm 1. Systems-learning co-design execution per communication round.")
    rr.bold=True; rr.font.name="Times New Roman"; rr.font.size=Pt(9.5)
    tbl=doc.add_table(rows=1+len(ALGORITHM1_ROWS), cols=2)
    tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
    tbl.style="Table Grid"
    tbl.autofit=False
    tbl.columns[0].width=Inches(0.42)
    tbl.columns[1].width=Inches(5.78)
    set_cell_text(tbl.cell(0,0),"Step",bold=True,color="FFFFFF",size=8.5)
    set_cell_text(tbl.cell(0,1),"Operation",bold=True,color="FFFFFF",size=8.5)
    set_cell_shading(tbl.cell(0,0),"17365D"); set_cell_shading(tbl.cell(0,1),"17365D")
    for ri,(step,op) in enumerate(ALGORITHM1_ROWS, start=1):
        set_cell_text(tbl.cell(ri,0),step,bold=True,size=8.2)
        set_cell_text(tbl.cell(ri,1),op,size=8.2)
        tbl.cell(ri,0).vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tbl.cell(ri,1).vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tbl.cell(ri,0).paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
        tbl.cell(ri,1).paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph()

def build_docx():
    doc=Document(); configure_doc(doc); lines=SOURCE.splitlines(); i=0; front_limit=lines.index("## Highlights")
    while i<len(lines):
        line=lines[i].rstrip()
        if i==0 and line.startswith("# "):
            p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(line[2:]).bold=True; i+=1; continue
        if line.startswith("## "): doc.add_paragraph(line[3:],style="Heading 1"); i+=1; continue
        if line.startswith("### "): doc.add_paragraph(line[4:],style="Heading 2"); i+=1; continue
        if line=="[[ALGORITHM:1]]":
            add_algorithm_box(doc); i+=1; continue
        if line.startswith("[[FIGURE:"):
            idx=int(re.search(r"(\d+)",line).group(1)); rel,cap=FIGURES[idx]
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(ROOT/rel),width=Inches(6.25))
            cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=cp.add_run(cap); rr.font.name="Times New Roman"; rr.font.size=Pt(9); rr.italic=True
            i+=1; continue
        if line.startswith("- "):
            p=doc.add_paragraph(style="List Bullet"); parse_inline_runs(p,line[2:]); i+=1; continue
        if re.match(r"^\*\*Table \d+\.", line):
            add_table_caption(doc, line); i+=1; continue
        if line.startswith("|") and i+1<len(lines) and re.match(r"^\|?\s*:?-+",lines[i+1]):
            block=[line]; i+=1
            while i<len(lines) and lines[i].startswith("|"): block.append(lines[i].rstrip()); i+=1
            rows=[]
            for k,row in enumerate(block):
                if k==1: continue
                rows.append([v.strip() for v in row.strip("|").split("|")])
            cols=max(len(r) for r in rows); tbl=doc.add_table(rows=len(rows),cols=cols); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER; tbl.style="Table Grid"
            for ri,row in enumerate(rows):
                for ci in range(cols):
                    txt=re.sub(r"\*\*","",row[ci] if ci<len(row) else "")
                    set_cell_text(tbl.cell(ri,ci),txt,bold=(ri==0),color="FFFFFF" if ri==0 else None)
                    tbl.cell(ri,ci).vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    if ri==0: set_cell_shading(tbl.cell(ri,ci),"17365D")
            doc.add_paragraph(); continue
        if line=="\\[":
            eq=[]; i+=1
            while i<len(lines) and lines[i].strip()!="\\]": eq.append(lines[i]); i+=1
            i+=1; add_math_para(doc," ".join(eq)); continue
        if not line: i+=1; continue
        p=doc.add_paragraph(); parse_inline_runs(p,line)
        if i<front_limit:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        i+=1
    out=OUT/"NEXUS_FINAL_MANUSCRIPT.docx"; doc.save(out); return out
def esc_tex(s):
    for a,b in [("&",r"\&"),("%",r"\%"),("#",r"\#"),("_",r"\_")]: s=s.replace(a,b)
    return s

def markdown_to_tex():
    lines=SOURCE.splitlines(); out=[
        r"\documentclass[11pt]{article}",r"\usepackage[a4paper,margin=2.2cm]{geometry}",r"\usepackage{graphicx}",
        r"\usepackage{amsmath,amssymb}",r"\usepackage{booktabs,longtable,array}",r"\usepackage{hyperref}",
        r"\usepackage{caption}",r"\usepackage{microtype}",r"\usepackage{xcolor}",r"\usepackage{enumitem}",
        r"\setlist{nosep}",r"\hypersetup{colorlinks=true,linkcolor=black,urlcolor=blue,citecolor=black}",r"\begin{document}"
    ]; i=0
    while i<len(lines):
        line=lines[i].rstrip()
        if i==0 and line.startswith("# "):
            out += [r"\begin{center}",r"{\LARGE\bfseries "+esc_tex(line[2:])+r"}\par\vspace{0.8em}",
                    r"{\large Abhinandan Tripathi$^{1}$, Vijay Kumar Tiwari$^{2}$, Mohd. Arif$^{3}$, Abhishek Kumar Pandey$^{4,*}$, Shivam Bhardwaj$^{5}$}\par",
                    r"$^{1}$Department of Computer Science and Engineering, Buddha Institute of Technology, Gorakhpur, India\par",
                    r"$^{2}$Department of Information Technology, Madan Mohan Malaviya University of Technology (MMMUT), Gorakhpur, India\par",
                    r"$^{3}$Department of Computer Science and Engineering, Galgotias University, Greater Noida, India\par",
                    r"$^{4}$Department of Computer Science and Engineering, Bennett University, Greater Noida, India\par",
                    r"$^{5}$United Institute of Management, Prayagraj, India\par",
                    r"Emails: $^{1}$\href{mailto:abhinandan282@bit.ac.in}{abhinandan282@bit.ac.in}; $^{2}$\href{mailto:vktitca@mmmut.ac.in}{vktitca@mmmut.ac.in}; $^{3}$\href{mailto:md.arif@galgotiasuniversity.edu.in}{md.arif@galgotiasuniversity.edu.in};\par",
                    r"$^{4}$\href{mailto:abhishek.pandey2@bennett.edu.in}{abhishek.pandey2@bennett.edu.in}; $^{5}$\href{mailto:shibambhardwaj@gmail.com}{shibambhardwaj@gmail.com}\par",
                    r"ORCID: Abhishek Kumar Pandey — 0000-0003-3799-9754; Shivam Bhardwaj — 0009-0005-4554-7397\par",
                    r"$^*$Corresponding author: Abhishek Kumar Pandey (\href{mailto:abhishek.pandey2@bennett.edu.in}{abhishek.pandey2@bennett.edu.in})\par",
                    r"All authors contributed equally to this work.",r"\end{center}\vspace{0.8em}"]
            i+=1
            while i<len(lines) and not lines[i].startswith("## Highlights"): i+=1
            continue
        if line.startswith("## "):
            h=line[3:]
            out.append(r"\section*{"+esc_tex(h)+r"}" if h in ("Highlights","In brief","Broader context","Resource availability","References","Abstract") else r"\section{"+esc_tex(re.sub(r"^\d+\.\s*","",h))+r"}")
            i+=1; continue
        if line.startswith("### "): out.append(r"\subsection*{"+esc_tex(re.sub(r"^\d+\.\d+\s*","",line[4:]))+r"}"); i+=1; continue
        if line=="[[ALGORITHM:1]]":
            out += [r"\begin{center}",r"\fbox{\begin{minipage}{0.95\linewidth}",r"\textbf{Algorithm 1. Systems--learning co-design execution per communication round.}",r"\begin{enumerate}"]
            for _,op in ALGORITHM1_ROWS:
                out.append(r"\item "+esc_tex(op))
            out += [r"\end{enumerate}",r"\end{minipage}}",r"\end{center}"]
            i+=1; continue
        if line.startswith("[[FIGURE:"):
            idx=int(re.search(r"(\d+)",line).group(1)); rel,cap=FIGURES[idx]; fn=Path(rel).name
            out += [r"\begin{figure}[htbp]",r"\centering",r"\includegraphics[width=0.96\linewidth]{figures/"+fn+r"}",r"\caption{"+esc_tex(re.sub(r"^Figure \d+\.\s*","",cap))+r"}",r"\end{figure}"]; i+=1; continue
        if line.startswith("- "):
            out.append(r"\begin{itemize}"); 
            while i<len(lines) and lines[i].startswith("- "):
                out.append(r"\item "+esc_tex(lines[i][2:].replace("**",""))); i+=1
            out.append(r"\end{itemize}"); continue
        if line.startswith("|") and i+1<len(lines) and re.match(r"^\|?\s*:?-+",lines[i+1]):
            block=[line]; i+=1
            while i<len(lines) and lines[i].startswith("|"): block.append(lines[i].rstrip()); i+=1
            rows=[]
            for k,row in enumerate(block):
                if k==1: continue
                rows.append([re.sub(r"\*\*","",v.strip()) for v in row.strip("|").split("|")])
            n=len(rows[0]); out += [r"\begin{table}[htbp]\centering\scriptsize",r"\resizebox{\linewidth}{!}{%",r"\begin{tabular}{l"+"r"*(n-1)+r"}\toprule"]
            for ri,row in enumerate(rows):
                out.append(" & ".join(esc_tex(v.replace("↓","").replace("↑","")) for v in row)+r" \\")
                if ri==0: out.append(r"\midrule")
            out += [r"\bottomrule\end{tabular}}",r"\end{table}"]; continue
        if line=="\\[":
            eq=[]; i+=1
            while i<len(lines) and lines[i].strip()!="\\]": eq.append(lines[i]); i+=1
            i+=1; out += [r"\["," ".join(eq),r"\]"]; continue
        if not line: i+=1; continue
        text=line.replace("**","").replace("*","")
        if re.match(r"^\[\d+\]",text): out.append(r"\noindent "+text.replace("&",r"\&")+r"\par"); i+=1; continue
        out.append(esc_tex(text)+r"\par"); i+=1
    out.append(r"\end{document}")
    return "\n".join(out)

def build_latex():
    for rel,_ in FIGURES: shutil.copy2(ROOT/rel,LATEX/"figures"/Path(rel).name)
    tex=markdown_to_tex(); (LATEX/"main.tex").write_text(tex); shutil.copy2(OUT/"NEXUS_MANUSCRIPT_SOURCE.md",LATEX/"NEXUS_MANUSCRIPT_SOURCE.md")
    return LATEX/"main.tex"

for stale in ("main.aux","main.log","main.out","main.pdf"):
    (LATEX/stale).unlink(missing_ok=True)
docx=build_docx(); tex=build_latex()
extras=OUT/"submission_extras"; extras.mkdir(exist_ok=True)
(extras/"highlights.txt").write_text("\n".join("• "+x for x in HIGHLIGHTS)+"\n")
(extras/"in_brief.txt").write_text(IN_BRIEF+"\n")
(extras/"broader_context.txt").write_text(BROADER+"\n")
zip_path=OUT/"NEXUS_LATEX_PROJECT.zip"
with zipfile.ZipFile(zip_path,"w",zipfile.ZIP_DEFLATED) as z:
    for p in LATEX.rglob("*"):
        if p.is_file() and p.suffix.lower() not in {".aux",".log",".out",".pdf"}: z.write(p,p.relative_to(LATEX))
print("DOCX",docx); print("TEX",tex); print("ZIP",zip_path)
