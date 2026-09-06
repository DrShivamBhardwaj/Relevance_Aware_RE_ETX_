from docx import Document
from docx.shared import Pt
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
DOCX = SCRIPT_DIR / "Sadhana_Relevance_Aware_RE_ETX_VALIDATION_UPDATED.docx"

BACKUP = SCRIPT_DIR / "Sadhana_Relevance_Aware_RE_ETX_BEFORE_VALIDATION_TEXT_UPDATE.docx"

if not BACKUP.exists():
    BACKUP.write_bytes(DOCX.read_bytes())

doc = Document(DOCX)


def find_para(startswith):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise RuntimeError(f"Paragraph not found: {startswith}")


def replace_para(startswith, new_text):
    p = find_para(startswith)
    p.text = new_text
    return p


def add_before(anchor, text, style="Body Text"):
    p = anchor.insert_paragraph_before(text)
    try:
        p.style = style
    except Exception:
        pass
    return p


def add_caption_before(anchor, text):
    return add_before(anchor, text, "Caption")


def add_table_before(anchor, rows):
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)

    # The Sadhana template does not contain Word's built-in
    # "Table Grid" style. Preserve the template's default
    # table formatting instead of forcing a missing style.
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            table.cell(i, j).text = str(value)

    # Move table from document end to immediately before anchor.
    anchor._p.addprevious(table._tbl)

    # Keep table text compact.
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8.5)

    return table


# ------------------------------------------------------------------
# ABSTRACT
# ------------------------------------------------------------------

replace_para(
    "Fixed-length reporting can spend",
    (
        "Fixed-length reporting can spend a substantial communication budget "
        "on routine observations in battery-constrained wireless sensor "
        "networks. This study evaluates a temporal-spatial relevance controller "
        "that adapts application payload length over residual-energy "
        "expected-transmission-count (RE-ETX) routing. A seven-configuration, "
        "ten-seed control study separates full-payload routing, equal-budget "
        "fixed reporting, shuffled relevance, temporal-only relevance and "
        "temporal-spatial relevance. Relative to fixed 4000-bit RE-ETX (C1), "
        "the retained temporal-spatial controller (C3) increases mean first-, "
        "half- and last-node death from 125.5, 409.1 and 1268.6 rounds to "
        "275.2, 870.1 and 2636.4 rounds, raises packet delivery ratio by "
        "2.462 percentage points, and reduces mean delay and communication "
        "energy per delivered report by approximately 54%. The equal-budget "
        "fixed control C5 uses 1744 bits versus 1747.8 attempted bits for C3 "
        "and reproduces most network-level gains; after Holm correction C3 "
        "retains advantages in HND, delay and energy per delivered report, "
        "whereas FND, LND, PDR and retries are not significantly different. "
        "The shuffled-relevance control C6 is statistically indistinguishable "
        "from C3 on the tested network outcomes after correction. In contrast, "
        "offline event-labelled evaluation shows a mean within-round "
        "event/non-event payload gap of 1218.7 bits for C3, compared with "
        "0 bits for C5 and 13.3 bits for C6; event payload preservation is "
        "0.802 for C3 versus 0.436 and 0.592, respectively. Calibrated "
        "topology tests spanning approximately 1.5 to 4.9 initial route hops "
        "retain the large C3-versus-C1 traffic-reduction benefit, while "
        "equal-budget C3-versus-C5 network differences become mixed in deeper "
        "multi-hop regimes. The evidence therefore supports relevance as a "
        "mechanism for directing payload budget toward event-active reports, "
        "but not a universal claim that relevance allocation itself improves "
        "all network metrics at an equal average bit budget."
    ),
)


# ------------------------------------------------------------------
# INTRODUCTION
# ------------------------------------------------------------------

replace_para(
    "This paper investigates a compact relevance-aware",
    (
        "This paper investigates a compact relevance-aware payload controller "
        "over an interpretable RE-ETX substrate. Temporal innovation and "
        "spatial disagreement produce a scalar score, and a safeguarded "
        "mapping converts that score to a report length. Reliability and "
        "residual energy determine forwarding costs independently. A separate "
        "relevance-age route term is retained only as an experimental factor. "
        "The expanded evaluation uses fixed-budget, shuffled-relevance and "
        "temporal-only controls so that traffic-volume reduction can be "
        "distinguished from source-to-relevance alignment."
    ),
)

replace_para(
    "The study makes three contributions.",
    (
        "The study makes four contributions. First, it specifies the "
        "implemented relevance, payload, routing and frame-accounting rules "
        "sufficiently to reproduce their interaction. Second, it introduces "
        "equal-budget fixed and shuffled-relevance controls that separate the "
        "network effect of transmitting fewer bits from the informational "
        "effect of assigning more bits to more relevant observations. Third, "
        "it evaluates relevance targeting against offline event truth using "
        "event payload preservation and a within-round event/non-event payload "
        "gap that removes round-level allocation confounding. Fourth, it tests "
        "the principal C1, C3 and equal-budget C5 configurations across "
        "calibrated shallow, medium and deeper multi-hop regimes in addition "
        "to the residual-energy-weight sweep."
    ),
)

replace_para(
    "The research hypothesis is deliberately limited:",
    (
        "The research hypotheses are deliberately separated. The first is "
        "that reducing the application payload under the stated radio model "
        "should improve reporting opportunities and communication service "
        "relative to fixed 4000-bit reporting. The second is that a genuine "
        "relevance controller should assign substantially more payload to "
        "event-active than to non-event reports when compared with an "
        "equal-budget fixed controller and a shuffled-relevance controller. "
        "A universal network-level advantage of relevance over every "
        "equal-budget allocation is not assumed; it is tested empirically."
    ),
)


# ------------------------------------------------------------------
# TABLE 3 — EXPANDED CONTROL DEFINITIONS
# ------------------------------------------------------------------

cap3 = find_para("Table 3. Configuration definitions")
cap3.text = "Table 3. Configuration definitions and experimental control roles."

t3 = doc.tables[2]

rows3 = [
    [
        "Paper case",
        "Age route term",
        "Payload/controller",
        "Experimental role",
    ],
    [
        "C1",
        "Off",
        "Fixed 4000 bits",
        "Full-payload RE-ETX baseline",
    ],
    [
        "C2",
        "On",
        "Fixed 4000 bits",
        "Relevance-age routing ablation",
    ],
    [
        "C3, retained method",
        "Off",
        "Temporal-spatial relevance adaptive",
        "Retained relevance-aware source controller",
    ],
    [
        "C4",
        "On",
        "Temporal-spatial relevance adaptive",
        "Combined payload + route-age ablation",
    ],
    [
        "C5",
        "Off",
        "Fixed 1744 bits",
        "Equal-budget source control",
    ],
    [
        "C6",
        "Off",
        "Round-wise shuffled relevance payloads",
        "Destroys source/relevance alignment at similar budget",
    ],
    [
        "C7",
        "Off",
        "Temporal-only relevance adaptive",
        "Spatial-feature ablation",
    ],
]

while len(t3.rows) < len(rows3):
    t3.add_row()

for i, row in enumerate(rows3):
    for j, value in enumerate(row):
        t3.cell(i, j).text = value


# ------------------------------------------------------------------
# EXPERIMENTAL DESIGN
# ------------------------------------------------------------------

replace_para(
    "The evaluation uses seeds {11, 23, 42",
    (
        "The evaluation uses seeds {11, 23, 42, 67, 101, 137, 173, 211, "
        "257, 307}. Each seed fixes topology, shadowing and sensing streams "
        "across configurations. Frame-error streams use the same seed but are "
        "consumed differently when payloads or routes change; individual "
        "channel outcomes are therefore not aligned counterfactually. The "
        "expanded main control study contains seventy runs: C1-C7 over all "
        "ten seeds. C5 fixes every attempted report at 1744 bits, the nearest "
        "lower byte-aligned value to the baseline-topology C3 attempted-report "
        "mean of 1747.84 bits. C6 preserves the current round's adaptive "
        "payload distribution but randomly permutes those payload assignments "
        "among sources. C7 sets the spatial relevance weight to zero so that "
        "only temporal innovation determines the adaptive payload. A separate "
        "forty-run sweep evaluates residual-energy weights 0, 0.25, 0.50 and "
        "1.00 on the same seed ensemble."
    ),
)

replace_para(
    "Table 3 resolves the archive",
    (
        "Table 3 separates the original four-case ablation from the three "
        "source-control extensions. Independent C1 and C3 executions at seed "
        "42 reproduce the pre-extension network trajectories, and repeated C6 "
        "execution reproduces the shuffled-control result exactly. Event truth "
        "is queried only after the relevance score has been computed and is "
        "used solely for offline evaluation; it is never supplied to payload "
        "selection or routing. The temporal-only C7 implementation bypasses "
        "the spatial feature calculation rather than merely setting its final "
        "coefficient to zero."
    ),
)

replace_para(
    "Two payload statistics require separate denominators.",
    (
        "Two payload statistics require separate denominators. The printed "
        "mean payload averages every original sensor position in every "
        "simulated round, including dead sensors. The attempted-report counter "
        "is therefore used for equal-budget controls. In the baseline topology "
        "the pooled C3 attempted-report mean is 1747.84 bits, or a 56.30% "
        "reduction from 4000 bits; C5 uses the byte-aligned 1744-bit fixed "
        "control. The archive's 'semantic radio efficiency' divides offered "
        "application bits by its radio-bit counter without restricting the "
        "numerator to delivered reports. The counter also reserves a full "
        "frame for a partial sender-death transmission. This ratio is retained "
        "only as a diagnostic and is not interpreted as recovered semantic "
        "information per physical bit."
    ),
)

replace_para(
    "Seeds, rather than individual reports",
    (
        "Seeds, rather than individual reports, are the experimental units. "
        "Descriptive results use the mean and sample standard deviation. For "
        "seed-wise differences, marginal 95% t intervals are reported. The "
        "original seven-outcome C1→C2, C1→C3 and C3→C4 analysis remains one "
        "21-test Holm family. A separate 21-test source-control family applies "
        "the same seven network outcomes to C1→C5, C5→C3 and C6→C3. The "
        "within-round targeting analysis uses the three planned contrasts "
        "C5→C3, C6→C3 and C3→C7 with Holm correction. Event-quality contrasts "
        "use separate six-metric Holm families. For topology sensitivity, "
        "eight outcomes are corrected within each topology/contrast family."
    ),
)


# ------------------------------------------------------------------
# NEW RESULTS SECTIONS 6.5–6.7
# ------------------------------------------------------------------

discussion = find_para("7. Discussion and limitations")

add_before(
    discussion,
    "6.5. Equal-budget and shuffled-relevance source controls",
    "Heading 2",
)

add_before(
    discussion,
    (
        "Table 6 compares the retained C3 controller with the fixed "
        "equal-budget C5 control, the shuffled-relevance C6 control and the "
        "temporal-only C7 ablation. C5 differs from C3 by only 3.84 attempted "
        "bits on the pooled baseline mean (1744 versus 1747.84 bits). "
        "Consequently, it reproduces most of the large C1→C3 network gain. "
        "For C5→C3, Holm-corrected differences remain for HND (+26.1 rounds), "
        "mean delay (−11.36 ms) and energy per delivered report "
        "(−0.0127 mJ), whereas FND, LND, PDR and retries per generated report "
        "are not significant after correction."
    ),
    "First Paragraph",
)

add_caption_before(
    discussion,
    "Table 6. Source-control outcomes, means over ten matched seeds.",
)

add_table_before(
    discussion,
    [
        ["Outcome", "C3", "C5", "C6", "C7"],
        ["Attempted payload, bits", "1747.84", "1744.00", "1748.05", "1610.89"],
        ["FND, rounds", "275.2", "267.6", "295.8", "297.8"],
        ["HND, rounds", "870.1", "844.0", "869.6", "929.8"],
        ["LND, rounds", "2636.4", "2626.6", "2634.0", "2825.8"],
        ["PDR, %", "97.812", "97.758", "97.825", "97.956"],
        ["Delay, ms", "481.6", "493.0", "481.2", "448.7"],
        ["Energy/delivered report, mJ", "0.526", "0.538", "0.525", "0.489"],
        ["Retries/generated report", "0.0918", "0.0941", "0.0922", "0.0861"],
    ],
)

add_before(
    discussion,
    (
        "C6 is especially informative because it preserves approximately the "
        "same aggregate adaptive-payload budget while destroying the mapping "
        "between a source's relevance and its assigned payload. None of the "
        "seven tested C6→C3 network outcomes remains significant after Holm "
        "correction. Thus, the dominant network-level improvement is explained "
        "by reduced traffic volume; the baseline network metrics alone do not "
        "show that relevance-aware allocation is superior to a shuffled "
        "allocation at the same average budget. C7 has a smaller mean payload "
        "than C3, so its better lifetime and delay outcomes cannot be treated "
        "as evidence that temporal-only relevance is semantically superior."
    ),
    "Body Text",
)

add_before(
    discussion,
    "6.6. Offline event targeting and relevance alignment",
    "Heading 2",
)

add_before(
    discussion,
    (
        "Network-level equality between C3 and C6 does not imply that the "
        "payload assignments are informationally equivalent. Event truth, "
        "which is unavailable to the runtime controller, permits an offline "
        "audit of where the bit budget is placed. Because event-active rounds "
        "can have a different overall payload distribution from inactive "
        "rounds, the principal targeting statistic is the event-minus-"
        "non-event payload difference computed within each round and then "
        "averaged across evaluable event rounds."
    ),
    "First Paragraph",
)

add_caption_before(
    discussion,
    "Table 7. Offline event-targeting diagnostics, means over ten matched seeds.",
)

add_table_before(
    discussion,
    [
        ["Metric", "C3", "C5", "C6", "C7"],
        ["Event report delivery ratio", "0.9784", "0.9815", "0.9810", "0.9841"],
        ["Event mean payload, bits", "3208.0", "1744.0", "2367.7", "1863.9"],
        ["Non-event mean payload, bits", "1681.4", "1744.0", "1719.8", "1599.3"],
        ["Event payload preservation", "0.8020", "0.4360", "0.5919", "0.4660"],
        ["Event payload delivery ratio", "0.7840", "0.4280", "0.5803", "0.4580"],
        ["Within-round targeting gap, bits", "1218.7", "0.0", "13.3", "201.6"],
    ],
)

add_before(
    discussion,
    (
        "The mean within-round targeting gap is 1218.7 bits for C3, compared "
        "with 0 for C5 and 13.3 bits for C6. The paired C5→C3 and C6→C3 "
        "differences are both highly significant after Holm correction "
        "(adjusted p < 10^-9). C3 also preserves 80.2% of the full 4000-bit "
        "event payload budget, compared with 43.6% for C5 and 59.2% for C6, "
        "and delivers 78.4% of the full event payload budget versus 42.8% and "
        "58.0%, respectively. Event-report delivery probability itself is not "
        "higher for C3 than C5. These results support a precise claim: the "
        "relevance estimator directs substantially more application bits "
        "toward event-active reports, but this targeting should not be "
        "described as event-detection recall or receiver reconstruction "
        "accuracy."
    ),
    "Body Text",
)

add_before(
    discussion,
    "6.7. Sensitivity to multi-hop topology depth",
    "Heading 2",
)

add_before(
    discussion,
    (
        "A separate topology calibration increases field depth while retaining "
        "100 sensors and positioning the sink 50 m above the upper field "
        "boundary. Three high-connectivity regimes were selected with initial "
        "mean route depths of approximately 1.483, 2.910 and 4.929 hops. "
        "Within each regime, C1, C3 and a topology-specific equal-budget C5 "
        "control are evaluated over all ten seeds. The C5 payloads are 1744, "
        "1648 and 1568 bits in the shallow, medium and deep profiles, closely "
        "matching the corresponding C3 attempted means of 1745.5, 1649.2 and "
        "1570.8 bits."
    ),
    "First Paragraph",
)

add_caption_before(
    discussion,
    (
        "Table 8. Topology-depth sensitivity. Cells show C1 / C3 / C5 means "
        "over ten seeds."
    ),
)

add_table_before(
    discussion,
    [
        [
            "Profile (initial hops)",
            "HND, rounds",
            "LND, rounds",
            "PDR, %",
            "Delay, ms",
            "Energy/delivered, mJ",
        ],
        [
            "Shallow (1.483)",
            "200.7 / 383.3 / 357.8",
            "1290.4 / 2642.5 / 2663.8",
            "90.496 / 95.404 / 95.294",
            "958.1 / 442.3 / 452.8",
            "1.631 / 0.744 / 0.764",
        ],
        [
            "Medium (2.910)",
            "49.9 / 112.3 / 96.2",
            "1063.3 / 2138.4 / 2206.6",
            "86.443 / 93.765 / 93.404",
            "2170.4 / 959.5 / 1002.1",
            "4.917 / 2.105 / 2.209",
        ],
        [
            "Deep (4.929)",
            "29.1 / 74.7 / 65.5",
            "934.4 / 1952.7 / 2303.9",
            "86.343 / 94.160 / 94.299",
            "2512.2 / 1050.6 / 1018.6",
            "7.604 / 3.082 / 2.983",
        ],
    ],
)

add_before(
    discussion,
    (
        "C3 remains substantially better than full-payload C1 on HND, LND, "
        "PDR, delay and communication energy as route depth increases, showing "
        "that the benefit of reducing frame workload is not confined to the "
        "original shallow topology. The equal-budget C5→C3 comparison is more "
        "qualified. C3 has higher HND in all three profiles, but LND is not "
        "improved in the medium profile and is 351.2 rounds lower than C5 in "
        "the deep profile after correction. In the deep profile, C3 also has "
        "no corrected advantage in PDR, delay, energy per delivered report or "
        "retry rate. The within-round targeting gap nevertheless remains large "
        "for C3 across shallow, medium and deep profiles (approximately "
        "1297, 1249 and 1124 bits). Hence deeper multi-hop testing reinforces "
        "the distinction between relevance alignment and generic traffic "
        "reduction: targeting persists, whereas equal-budget network "
        "superiority is metric- and topology-dependent."
    ),
    "Body Text",
)


# ------------------------------------------------------------------
# DISCUSSION
# ------------------------------------------------------------------

replace_para(
    "The evidence supports source-length adaptation over RE-ETX",
    (
        "The expanded evidence separates two effects that were confounded in "
        "the original four-case analysis. First, transmitting fewer bits is "
        "responsible for most of the large network-level improvement relative "
        "to 4000-bit C1. The equal-budget C5 and shuffled C6 controls reproduce "
        "most of C3's lifetime, PDR, delay and energy behaviour. Second, the "
        "relevance estimator is not arbitrary: C3 assigns far more payload to "
        "event-active reports than either the fixed-budget or shuffled "
        "controls, with a within-round targeting gap above 1.2 kbit in the "
        "baseline topology. The defensible contribution is therefore "
        "relevance-directed allocation of a reduced payload budget, not a "
        "claim that relevance uniquely causes every network-level gain."
    ),
)

replace_para(
    "Several modelling decisions constrain external validity.",
    (
        "Several modelling decisions still constrain external validity. The "
        "original topology averages only about 1.4 delivered hops, but the "
        "new sensitivity study extends calibrated initial route depth to "
        "approximately 4.9 hops. The C3-versus-C1 advantage survives this "
        "extension, whereas equal-budget C3-versus-C5 outcomes become mixed "
        "and deep-profile LND favours C5. The sensitivity profiles change field "
        "height and candidate-link radius while holding node count fixed, so "
        "they should not be interpreted as a pure hop-count intervention or "
        "as a large-network scalability experiment. Static shadowing, "
        "independent frame errors and serialised access still exclude "
        "contention, interference, burst losses and realistic acknowledgements. "
        "The first-order radio model also omits measured processor and "
        "transceiver-state costs."
    ),
)


# ------------------------------------------------------------------
# CONCLUSIONS
# ------------------------------------------------------------------

replace_para(
    "A safeguarded temporal-spatial relevance controller has been evaluated",
    (
        "A safeguarded temporal-spatial relevance controller has been "
        "evaluated over residual-energy ETX routing using seven source/routing "
        "configurations and ten matched seeds. Relative to fixed 4000-bit C1, "
        "C3 approximately doubles the major sensing-round lifetime milestones "
        "and reduces delay and communication energy per delivered report by "
        "about 54%. Equal-budget and shuffled controls show, however, that "
        "most of these network-level gains arise from transmitting fewer bits. "
        "At approximately the same average payload, C3 and shuffled C6 are not "
        "significantly different on the seven tested network outcomes after "
        "Holm correction."
    ),
)

replace_para(
    "The practical interpretation is narrower",
    (
        "The relevance-specific evidence lies in payload placement. C3 "
        "produces a mean within-round event/non-event payload gap of "
        "1218.7 bits, versus 0 for the fixed equal-budget control and "
        "13.3 bits for shuffled relevance, while preserving about 80.2% of "
        "the full event payload budget. This targeting remains substantial "
        "under the deeper topology profiles, even though equal-budget network "
        "advantages are not universal and deep-profile LND favours C5. The "
        "result therefore supports relevance-aware payload prioritisation as "
        "an auditable traffic-allocation mechanism, while receiver fidelity, "
        "application utility, processing energy, neighbourhood-information "
        "exchange and larger-scale contention-aware deployments remain open "
        "validation requirements."
    ),
)


# ------------------------------------------------------------------
# DATA / CODE AVAILABILITY
# ------------------------------------------------------------------

replace_para(
    "The project repository is",
    (
        "The project repository is "
        "https://github.com/DrShivamBhardwaj/Relevance_Aware_RE_ETX_. "
        "The validated experimental package is available on branch "
        "'dual-journal-validation' and is anchored by tag "
        "'dual-journal-validation-v1'. The repository includes the C1-C7 "
        "multi-seed control data, event-quality and within-round targeting "
        "statistics, residual-energy-weight sweep, calibrated topology-depth "
        "sensitivity runs, paired tests, analysis scripts, batch runners, "
        "compressed per-round histories and reproducibility metadata."
    ),
)


doc.save(DOCX)

print("UPDATED MANUSCRIPT SAVED:")
print(DOCX)
print()
print("BACKUP PRESERVED:")
print(BACKUP)
print()
print("TABLE COUNT:", len(doc.tables))

for key in [
    "Fixed-length reporting can spend",
    "The study makes four contributions",
    "The evaluation uses seeds",
    "6.5. Equal-budget",
    "6.6. Offline event",
    "6.7. Sensitivity",
    "The expanded evidence separates",
    "A safeguarded temporal-spatial relevance controller",
]:
    for p in doc.paragraphs:
        if p.text.strip().startswith(key):
            print("FOUND:", p.text[:140])
            break
